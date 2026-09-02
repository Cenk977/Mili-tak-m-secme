"""
Export rankings to Excel format.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import logging

logger = logging.getLogger(__name__)


def create_rankings_xlsx(athletes: list, leg: str = 'combined') -> bytes:
    """
    Create Excel workbook with rankings.

    Args:
        athletes: List of athlete dicts from API
        leg: 'antalya', 'edirne', or 'combined'

    Returns:
        XLSX file as bytes
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sıralamalar"

    # Define styles
    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")

    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Add title row
    ws.merge_cells('A1:H1')
    title = ws['A1']
    title.value = f"Milli Takım Seçme - {leg.upper()} Sıralaması"
    title.font = Font(bold=True, size=14)
    title.alignment = Alignment(horizontal="center")

    # Add headers
    headers = ['Sıra', 'Adı Soyadı', 'Doğum Yılı', 'Cinsiyet', 'Şehir', 'Bölge', 'Puan', 'Seçim']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border

    # Add athlete rows
    for idx, athlete in enumerate(athletes, 1):
        row = idx + 3

        ws.cell(row=row, column=1).value = idx
        ws.cell(row=row, column=2).value = athlete.get('athlete_name', '')
        ws.cell(row=row, column=3).value = athlete.get('birth_year', '')
        ws.cell(row=row, column=4).value = athlete.get('gender', '')
        ws.cell(row=row, column=5).value = athlete.get('city', '')
        ws.cell(row=row, column=6).value = athlete.get('region', '')
        ws.cell(row=row, column=7).value = round(athlete.get('display_top3', 0), 2)
        ws.cell(row=row, column=8).value = athlete.get('selected', '-')

        # Apply border to all cells
        for col in range(1, 9):
            ws.cell(row=row, column=col).border = border

    # Adjust column widths
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 10
    ws.column_dimensions['H'].width = 12

    # Save to bytes
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    logger.info(f"Generated XLSX with {len(athletes)} athletes for leg {leg}")
    return output.getvalue()
